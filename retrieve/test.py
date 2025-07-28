from docx import Document
from pprint import pprint

# === 读取 Word 文档 ===
doc = Document("merged_table_example.docx")

# 分离段落和表格
paragraphs = []
tables = []

for block in doc.element.body:
    if block.tag.endswith("tbl"):
        tables.append(doc.tables[len(tables)])  # 用 doc.tables 顺序访问
    elif block.tag.endswith("p"):
        para = block.xpath(".//w:t")
        if para:
            text = ''.join([t.text for t in para if t.text])
            if text.strip():
                paragraphs.append(text.strip())

# === 处理第一个表格（合并行展开） ===
def read_and_fill_merged_table(docx_table):
    table_data = []
    for row in docx_table.rows:
        table_data.append([cell.text.strip() for cell in row.cells])
    print(table_data)

    return table_data

filled_table = read_and_fill_merged_table(tables[0])
print(filled_table)