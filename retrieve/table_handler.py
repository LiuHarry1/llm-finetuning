from docx import Document
from pprint import pprint

# === 读取 Word 文档 ===
doc = Document("./data/test_table.docx")

# 分离段落和表格
paragraphs = []
tables = []

for block in doc.element.body:
    if block.tag.endswith("tbl"):
        tables.append(doc.tables[len(tables)])  # 用 doc.tables 顺序访问
    elif block.tag.endswith("p"):
        para = block.xpath(".//w:t")
        if para:
            text = ''.join([t.text for t in para if t.text])
            if text.strip():
                paragraphs.append(text.strip())

# === 处理第一个表格（合并行展开） ===
def read_and_fill_merged_table(docx_table):
    table_data = []
    for row in docx_table.rows:
        table_data.append([cell.text.strip() for cell in row.cells])
    print(table_data)
    # 填补合并单元格（按行方向）
    headers = table_data[0]
    filled = [headers]
    last_values = [""] * len(headers)
    for row in table_data[1:]:
        new_row = []
        for i, cell in enumerate(row):
            if cell.strip():
                last_values[i] = cell
            new_row.append(last_values[i])
        filled.append(new_row)
    return filled

filled_table = read_and_fill_merged_table(tables[0])

# === 表格转 JSON ===
def table_to_json(filled_table):
    headers = filled_table[0]
    rows = filled_table[1:]
    return [dict(zip(headers, row)) for row in rows]

json_table = table_to_json(filled_table)

# === 表格转为自然语言 ===
def table_to_sentences(rows):
    sentences = []
    for row in rows:
        sentence = (
            f"{row['类别']}类产品{row['产品']}的销售额为{row['销售额']}万元，"
            f"利润率为{row['利润率']}。"
        )
        sentences.append(sentence)
    return sentences

table_sentences = table_to_sentences(json_table)

# === 拼接最终文本 ===
final_text = "\n\n".join(paragraphs[:2] + table_sentences + paragraphs[2:])

# === 输出 ===
print("✅ 提取表格为 JSON：")
pprint(json_table)

print("\n✅ 拼接后文本（适合 RAG 向量化）：")
print(final_text)
