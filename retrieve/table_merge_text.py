from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_vertical_merge(cell, merge_type="restart"):
    """设置垂直合并属性"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vMerge = OxmlElement("w:vMerge")
    vMerge.set(qn("w:val"), merge_type)
    tcPr.append(vMerge)

def set_colspan(cell, span_val):
    """设置水平合并属性"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    gridSpan = OxmlElement("w:gridSpan")
    gridSpan.set(qn("w:val"), str(span_val))
    tcPr.append(gridSpan)

# 创建 Word 文档
doc = Document()
doc.add_heading("简单示例文档", level=1)
doc.add_paragraph("这是一个包含合并单元格的表格，用于测试结构解析。")

# 创建 3x3 表格
table = doc.add_table(rows=3, cols=3)
table.style = 'Table Grid'

# 第 1 行：合并前两列
cell_00 = table.cell(0, 0)
cell_00.text = "产品信息"
set_colspan(cell_00, 2)
table.cell(0, 2).text = "利润率"

# 第 2 行
table.cell(1, 0).text = "饮料"
set_vertical_merge(table.cell(1, 0), "restart")
table.cell(1, 1).text = "可乐"
table.cell(1, 2).text = "20%"

# 第 3 行
set_vertical_merge(table.cell(2, 0), "continue")
table.cell(2, 1).text = "果汁"
table.cell(2, 2).text = "18%"

# 保存
output_path = "simple_merged_table.docx"
doc.save(output_path)
print(f"✅ 已生成 Word 文件：{output_path}")
